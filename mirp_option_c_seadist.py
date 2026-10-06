"""
Stages 1 and 2 of the three-stage MIRP framework
=======================================================================
Maritime Inventory Routing with Time-Charter Fleet Sizing and
Supplementary Voyage Chartering for Bulk Cement Distribution
(Maritime Policy & Management manuscript; see README.md).

Stage 1  K-Means clustering of the nine destination ports (k-means++,
         20 restarts, seed 42, Euclidean distance on latitude/longitude),
         then the visit sequence of each cluster by complete enumeration
         over the maritime sailing-distance matrix (data/sea_distance_
         matrix_nm.csv). The matrix is the only sailing-distance authority.
Stage 2  Time-indexed MILP (2-day periods, T_MAX = 183, i.e. 366 modeled
         days) that selects the committed time-charter fleet endogenously
         and jointly determines cluster assignment, voyage scheduling,
         cargo allocation and destination-silo inventory, with unmet
         demand penalized (soft constraint).  Cargo-space utilization of
         the deployed time-charter voyages is a reporting quantity only.

Stage 3 (supplementary voyage chartering) is in
supplementary_charter_milp_seadist_sweep.py.

Run one configuration:   K_CLUSTERS=5 python mirp_option_c_seadist.py
Run K = 2..9:            python run_seadist_ksweep.py
Solver: Gurobi (the reported runs used Gurobi 12.0.1; OR-Tools fallback is
provided for convenience only and was not used for the reported results).
Vessels are identified only by anonymized codes CC1 to CC8.
"""

import math
import csv
import os
import random
import itertools
import time as _time
import hashlib                  # provenance checksums
import json                     # distance metadata + run manifest
import platform                 # run manifest (hardware/platform)
import sys                      # run manifest (Python version)
import datetime as _dt          # run manifest timestamp
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

# ── Solver import ──
try:
    import gurobipy as gp
    from gurobipy import GRB
    USE_GUROBI = True
    print("[Solver] Gurobi detected")
except ImportError:
    USE_GUROBI = False
    try:
        from ortools.linear_solver import pywraplp
        print("[Solver] Gurobi not found, using OR-Tools fallback")
    except ImportError:
        raise ImportError("Neither gurobipy nor ortools installed. "
                          "Install one: pip install gurobipy OR pip install ortools")


# =============================================================
# GLOBAL PARAMETERS
# =============================================================
PERIOD_DAYS     = 2             # Days per time period
TOTAL_DAYS      = 365           # Planning horizon (days)
T_MAX           = math.ceil(TOTAL_DAYS / PERIOD_DAYS)  # = 183 periods
COMMISSION_DAYS = 350           # Operating days per year, following Stopford (2009)
KURS            = 15_857        # Rp per US dollar (KURS); enters the time-charter hire, Eq. (4)
HARGA_IFO       = 18_600        # IFO fuel price (Rp per liter; fuel volumes are computed in liters)
HARGA_MDO       = 24_000        # MDO fuel price (Rp per liter)
CHC_RATE        = 700           # Cargo handling cost (Rp/ton)
STOCKOUT_PENALTY = 1_225_000    # Stockout penalty = cement selling price (Rp/ton)
BUFFER_TIME_PCT = 0.10          # 10% buffer for non-productive time
K_CLUSTERS      = int(os.environ.get('K_CLUSTERS', '5'))  # Number of K-Means clusters (reported runs: K = 2..9)
RANDOM_SEED     = 42
TIME_LIMIT_SEC  = 7200          # Solver time limit (120 minutes)


# =============================================================
# PORT DATA
# =============================================================
ports = {
    'Biringkassi': {
        'role': 'origin',
        'loading_rate': 1500,  # ton/hour
    },
    'Bitung': {
        'demand': 69822, 'capacity': 12000, 'depth': 10.5,
        'unloading_rate': 300, 'initial_stock': 1900,
        'holding_cost': 705.48,
    },
    'Mamuju': {
        'demand': 9021, 'capacity': 4000, 'depth': 10.0,
        'unloading_rate': 300, 'initial_stock': 400,
        'holding_cost': 705.48,
    },
    'Palu': {
        'demand': 338083, 'capacity': 8000, 'depth': 10.0,
        'unloading_rate': 300, 'initial_stock': 300,
        'holding_cost': 705.48,
    },
    'Kendari': {
        'demand': 145158, 'capacity': 12000, 'depth': 10.0,
        'unloading_rate': 300, 'initial_stock': 2000,
        'holding_cost': 705.48,
    },
    'Ambon': {
        'demand': 12205, 'capacity': 8000, 'depth': 9.2,
        'unloading_rate': 300, 'initial_stock': 2000,
        'holding_cost': 624.66,
    },
    'Oba': {
        'demand': 80930, 'capacity': 6000, 'depth': 9.0,
        'unloading_rate': 300, 'initial_stock': 725,
        'holding_cost': 624.66,
    },
    'Celukan Bawang': {
        'demand': 129014, 'capacity': 12000, 'depth': 9.6,
        'unloading_rate': 100, 'initial_stock': 1000,
        'holding_cost': 621.92,
    },
    'Lembar': {
        'demand': 61222, 'capacity': 5000, 'depth': 11.5,
        'unloading_rate': 300, 'initial_stock': 600,
        'holding_cost': 616.44,
    },
    'Sorong': {
        'demand': 50551, 'capacity': 12000, 'depth': 11.3,
        'unloading_rate': 300, 'initial_stock': 2200,
        'holding_cost': 726.03,
    },
}

dest_port_names = [p for p in ports if p != 'Biringkassi']

# Port coordinates (decimal degrees) are used ONLY by the Stage-1 K-Means.
# They are not distributed with this repository (see README, Data availability).
# If data/port_coordinates.csv (columns: port,lat,lon) is supplied, K-Means is
# re-executed and its memberships are checked against RECORDED_KMEANS_CLUSTERS;
# otherwise the recorded Stage-1 memberships of the reported runs are used.
PORT_COORDINATES_CSV = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                    'data', 'port_coordinates.csv')


def _load_port_coordinates():
    if not os.path.exists(PORT_COORDINATES_CSV):
        return False
    with open(PORT_COORDINATES_CSV, newline='', encoding='utf-8') as fh:
        rows = {r['port']: r for r in csv.DictReader(fh)}
    missing = [p for p in ports if p not in rows]
    if missing:
        raise SystemExit(f"[FATAL] port_coordinates.csv lacks ports: {missing}")
    for p in ports:
        ports[p]['lat'] = float(rows[p]['lat'])
        ports[p]['lon'] = float(rows[p]['lon'])
    return True


HAS_PORT_COORDINATES = _load_port_coordinates()


# =============================================================
# VESSEL DATA: 8 candidate vessels, anonymized codes CC1 to CC8
# =============================================================
ships = [
    {'id': 'CC1', 'name': 'CC1',              'payload': 7000,  'GT': 4676,  'DWT': 7325,
     'HP_ME': 3746.754,  'HP_AE': 936.6885,   'draft': 7.23, 'vs_b': 11.0, 'vs_l': 8.4,  'prod_bm': 325},
    {'id': 'CC2', 'name': 'CC2',              'payload': 4000,  'GT': 2472,  'DWT': 4399,
     'HP_ME': 3549.627,  'HP_AE': 887.40675,  'draft': 6.9,  'vs_b': 11.3, 'vs_l': 7.9,  'prod_bm': 300},
    {'id': 'CC3', 'name': 'CC3',              'payload': 3600,  'GT': 2294,  'DWT': 3970,
     'HP_ME': 3549.627,  'HP_AE': 887.40675,  'draft': 4.8,  'vs_b': 12.9, 'vs_l': 6.7,  'prod_bm': 250},
    {'id': 'CC4', 'name': 'CC4',              'payload': 3400,  'GT': 3779,  'DWT': 3683,
     'HP_ME': 2958.246,  'HP_AE': 739.5615,   'draft': 4.4,  'vs_b': 11.8, 'vs_l': 8.0,  'prod_bm': 250},
    {'id': 'CC5', 'name': 'CC5',              'payload': 3400,  'GT': 2135,  'DWT': 3745,
     'HP_ME': 2716.866,  'HP_AE': 679.2165,   'draft': 4.6,  'vs_b': 12.2, 'vs_l': 5.9,  'prod_bm': 250},
    {'id': 'CC6', 'name': 'CC6',              'payload': 4000,  'GT': 2637,  'DWT': 4486,
     'HP_ME': 3746.754,  'HP_AE': 936.6885,   'draft': 6.89, 'vs_b': 12.5, 'vs_l': 9.0,  'prod_bm': 400},
    {'id': 'CC7', 'name': 'CC7',              'payload': 6000,  'GT': 3828,  'DWT': 6706,
     'HP_ME': 3549.627,  'HP_AE': 887.40675,  'draft': 5.8,  'vs_b': 12.5, 'vs_l': 8.2,  'prod_bm': 300},
    {'id': 'CC8', 'name': 'CC8',              'payload': 4500,  'GT': 3568,  'DWT': 5461,
     'HP_ME': 3648.861,  'HP_AE': 912.21525,  'draft': 5.7,  'vs_b': 17.9, 'vs_l': 11.6, 'prod_bm': 600},
]

N_SHIPS = len(ships)


# =============================================================
# UTILITY FUNCTIONS
# =============================================================

# =============================================================
# SAILING-DISTANCE AUTHORITY
# =============================================================
# The maritime sailing-distance matrix (nautical miles; Table A5 of the
# supplementary appendix), extracted from the confidential case workbook, is
# the ONLY sailing-distance authority in this script. Values are used EXACTLY
# AS RECORDED: no shortest-path closure, no averaging, no imputation.
_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
SEA_DISTANCE_CSV = os.path.join(_SCRIPT_DIR, 'data', 'sea_distance_matrix_nm.csv')
SEA_DISTANCE_META = os.path.join(_SCRIPT_DIR, 'data', 'sea_distance_matrix_nm.meta.json')
EXPECTED_SEA_CSV_SHA256 = 'fdc433907aeb90247ff5fa44add3f569cac7767d57a3437f4d3b5f4bfbc54dfe'
EXPECTED_SEA_SOURCE_SHA256 = '9a6a3b02cecdc3c512502960cfdc14c65de8f17d0ac96f87a45c5368a70a2b64'
SEA_CODE_TO_PORT = {
    'BRK': 'Biringkassi', 'BT': 'Bitung', 'MMJ': 'Mamuju', 'PL': 'Palu',
    'KDR': 'Kendari', 'AMB': 'Ambon', 'OB': 'Oba', 'LBR': 'Lembar',
    'SRG': 'Sorong', 'CB': 'Celukan Bawang',
}


def _sha256_file(path):
    h = hashlib.sha256()
    with open(path, 'rb') as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()


def _load_sea_distance_matrix(model_ports):
    """Load and validate the sea-distance matrix. Fails loudly on any defect."""
    for path in (SEA_DISTANCE_CSV, SEA_DISTANCE_META):
        if not os.path.exists(path):
            raise SystemExit(f"[FATAL] sea-distance input missing: {path}")
    csv_sha = _sha256_file(SEA_DISTANCE_CSV)
    if csv_sha != EXPECTED_SEA_CSV_SHA256:
        raise SystemExit(f"[FATAL] sea-distance CSV checksum {csv_sha} != expected "
                         f"{EXPECTED_SEA_CSV_SHA256}")
    with open(SEA_DISTANCE_META, encoding='utf-8') as fh:
        meta = json.load(fh)
    if meta.get('csv_sha256') != csv_sha:
        raise SystemExit("[FATAL] sea-distance meta csv_sha256 does not match the CSV")
    if meta.get('source_sha256') != EXPECTED_SEA_SOURCE_SHA256:
        raise SystemExit("[FATAL] sea-distance meta source_sha256 is not the recorded source workbook")
    if meta.get('unit') != 'nautical_miles':
        raise SystemExit(f"[FATAL] sea-distance unit is {meta.get('unit')!r}, expected nautical_miles")
    with open(SEA_DISTANCE_CSV, newline='', encoding='utf-8') as fh:
        rows = list(csv.reader(fh))
    header, body = rows[0], rows[1:]
    codes = header[2:]
    if header[:2] != ['code', 'port'] or sorted(codes) != sorted(SEA_CODE_TO_PORT) or len(codes) != 10:
        raise SystemExit(f"[FATAL] unexpected sea-distance CSV header: {header}")
    if [r[0] for r in body] != codes:
        raise SystemExit("[FATAL] sea-distance CSV row codes do not match column codes")
    names = [SEA_CODE_TO_PORT[c] for c in codes]
    for r, name in zip(body, names):
        if r[1] != name:
            raise SystemExit(f"[FATAL] sea-distance CSV port name {r[1]!r} != expected {name!r}")
    if set(names) != set(model_ports) or len(set(names)) != 10:
        raise SystemExit(f"[FATAL] sea-distance ports {sorted(names)} != model ports {sorted(model_ports)}")
    mat = {}
    for r, a in zip(body, names):
        for val, b in zip(r[2:], names):
            v = float(val)
            if not math.isfinite(v):
                raise SystemExit(f"[FATAL] non-finite sea distance {a}-{b}")
            mat[(a, b)] = v
    n_pairs = 0
    for i, a in enumerate(names):
        if mat[(a, a)] != 0.0:
            raise SystemExit(f"[FATAL] non-zero diagonal for {a}")
        for b in names[i + 1:]:
            if not (mat[(a, b)] > 0):
                raise SystemExit(f"[FATAL] non-positive sea distance {a}-{b}")
            if mat[(a, b)] != mat[(b, a)]:
                raise SystemExit(f"[FATAL] asymmetric sea distance {a}-{b}")
            n_pairs += 1
    if n_pairs != 45:
        raise SystemExit(f"[FATAL] expected 45 unordered port pairs, found {n_pairs}")
    return mat, meta, csv_sha


# Validated at import time, i.e. before any model can be built.
SEA_NM, SEA_META, SEA_CSV_SHA256 = _load_sea_distance_matrix(set(ports.keys()))
_SEA_PORTS = frozenset(ports.keys())


def sea_nm(origin_name, destination_name):
    """Authoritative sailing distance in nautical miles (maritime matrix)."""
    if origin_name not in _SEA_PORTS:
        raise KeyError(f"sea_nm: unknown port {origin_name!r}")
    if destination_name not in _SEA_PORTS:
        raise KeyError(f"sea_nm: unknown port {destination_name!r}")
    if origin_name == destination_name:
        return 0.0
    return SEA_NM[(origin_name, destination_name)]


def _environment_record():
    """Software/hardware record for run manifests."""
    rec = {
        'python_version': sys.version,
        'platform': platform.platform(),
        'processor': platform.processor() or platform.machine(),
        'logical_cpus': os.cpu_count(),
    }
    try:
        import psutil  # optional
        rec['ram_gb'] = round(psutil.virtual_memory().total / 1024 ** 3, 1)
    except Exception:
        rec['ram_gb'] = 'unavailable (psutil not installed)'
    try:
        rec['gurobi_version'] = '.'.join(str(x) for x in gp.gurobi.version()) if USE_GUROBI else None
    except Exception as exc:
        rec['gurobi_version'] = f'unavailable ({exc})'
    return rec


def _write_manifest(path, payload):
    """Write a JSON manifest; refuses to overwrite."""
    if os.path.exists(path):
        raise SystemExit(f"[FATAL] manifest already exists, refusing to overwrite: {path}")
    with open(path, 'w', encoding='utf-8') as fh:
        json.dump(payload, fh, indent=2, ensure_ascii=False, default=str)


def compute_biaya_tunda(GT):
    """Towage fee schedule based on GT (Indonesian port authority tariff)."""
    if GT <= 2000:    return 367_500
    elif GT <= 3500:  return 486_500
    elif GT <= 8000:  return 755_000
    elif GT <= 14000: return 1_171_000
    else:             return 1_500_000


# =============================================================
# STAGE 1a: K-MEANS CLUSTERING
# =============================================================

def kmeans_clustering(port_names, k, seed=RANDOM_SEED, n_init=20, max_iter=100):
    """K-Means clustering on destination ports based on (lat, lon)."""
    random.seed(seed)
    coords = [(ports[p]['lat'], ports[p]['lon']) for p in port_names]
    n = len(coords)

    def eucl(a, b):
        return math.sqrt((a[0]-b[0])**2 + (a[1]-b[1])**2)

    best_sse = float('inf')
    best_labels = None

    for _ in range(n_init):
        # K-means++ init
        centroids = [coords[random.randint(0, n-1)]]
        for _ in range(1, k):
            dists = [min(eucl(c, ct)**2 for ct in centroids) for c in coords]
            total = sum(dists)
            if total == 0:
                centroids.append(coords[random.randint(0, n-1)])
                continue
            r = random.random() * total
            cum = 0
            for idx, d in enumerate(dists):
                cum += d
                if cum >= r:
                    centroids.append(coords[idx])
                    break

        for _ in range(max_iter):
            labels = [min(range(k), key=lambda c: eucl(coords[i], centroids[c]))
                      for i in range(n)]
            new_centroids = []
            for c in range(k):
                members = [coords[i] for i in range(n) if labels[i] == c]
                if members:
                    new_centroids.append((
                        sum(m[0] for m in members) / len(members),
                        sum(m[1] for m in members) / len(members)))
                else:
                    new_centroids.append(centroids[c])
            if new_centroids == centroids:
                break
            centroids = new_centroids

        sse = sum(eucl(coords[i], centroids[labels[i]])**2 for i in range(n))
        if sse < best_sse:
            best_sse = sse
            best_labels = labels[:]

    clusters = {}
    for i, lbl in enumerate(best_labels):
        clusters.setdefault(lbl, []).append(port_names[i])
    return [clusters[c] for c in sorted(clusters.keys())], best_sse


# =============================================================
# STAGE 1b: TSP PER CLUSTER
# =============================================================

def solve_tsp(cluster_ports):
    """Optimal port visit sequence minimizing round-trip distance (NM)."""
    origin = ports['Biringkassi']
    if len(cluster_ports) == 1:
        return cluster_ports[:]

    best_dist = float('inf')
    best_order = None

    for perm in itertools.permutations(cluster_ports):
        total = 0
        prev_name = 'Biringkassi'
        for p_name in perm:
            total += sea_nm(prev_name, p_name)
            prev_name = p_name
        total += sea_nm(prev_name, 'Biringkassi')
        if total < best_dist:
            best_dist = total
            best_order = list(perm)

    return best_order


# =============================================================
# PRE-COMPUTATION: Arrival Offsets, RTD, Voyage Costs
# =============================================================

def precompute_vessel_cluster(ship, cluster_ports, route_order):
    """
    For a given ship and cluster route, compute:
      - feasibility (draft check)
      - arrival_offsets: {port_name: offset_in_periods} from departure
      - RTD_periods: round-trip duration in periods
      - voyage_cost: fuel + port charges per voyage
      - Various time components
    """
    origin = ports['Biringkassi']
    payload = ship['payload']
    D_c = sum(ports[p]['demand'] for p in cluster_ports)

    # Draft feasibility
    min_depth = min(ports[p]['depth'] for p in cluster_ports)
    if ship['draft'] > min_depth:
        return {'feasible': False, 'reason': 'draft'}

    # ── Timeline (in days) ──
    # Loading at Biringkassi
    load_days = payload / origin['loading_rate'] / 24

    # Build cumulative timeline
    cum_days = load_days
    arrival_offsets = {}   # port → arrival period offset
    unload_times = {}      # port → unload time in days

    prev_name = 'Biringkassi'
    total_sail_dist = 0

    for p_name in route_order:
        p = ports[p_name]
        dist = sea_nm(prev_name, p_name)
        total_sail_dist += dist
        sail_days = dist / (ship['vs_l'] * 24)
        cum_days += sail_days

        # Arrival period (offset from departure period)
        arrival_offsets[p_name] = max(1, math.ceil(cum_days / PERIOD_DAYS))

        # Unloading time (proportional allocation for timing estimate)
        share = ports[p_name]['demand'] / D_c
        cargo_est = payload * share
        ul_days = cargo_est / ports[p_name]['unloading_rate'] / 24
        unload_times[p_name] = ul_days
        cum_days += ul_days

        prev_name = p_name

    # Return sailing (ballast)
    return_dist = sea_nm(route_order[-1], 'Biringkassi')
    total_sail_dist += return_dist
    return_days = return_dist / (ship['vs_b'] * 24)
    cum_days += return_days

    # Buffer
    buffer_days = cum_days * BUFFER_TIME_PCT
    cum_days += buffer_days

    RTD_periods = max(2, math.ceil(cum_days / PERIOD_DAYS))
    RTD_days = cum_days

    # Max frequency
    max_freq = math.floor(COMMISSION_DAYS / RTD_days) if RTD_days > 0 else 0
    if max_freq < 1:
        return {'feasible': False, 'reason': 'RTD too long'}

    # ── Voyage Cost ──
    GT = ship['GT']
    HP_ME = ship['HP_ME']
    HP_AE = ship['HP_AE']
    sea_days = cum_days - load_days - sum(unload_times.values()) - buffer_days
    port_days = load_days + sum(unload_times.values()) + buffer_days

    # Fuel cost per voyage
    me_base = (HP_ME * 0.18 * 24) / 1000 * 1123.6
    me_per_voy = me_base * 1.0 * sea_days + me_base * 0.2 * port_days
    ae_base = (HP_AE * 0.18 * 24) / 1000 * 1201.92
    ae_per_voy = ae_base * 0.5 * sea_days + ae_base * 1.0 * port_days
    fuel_per_voy = me_per_voy * HARGA_IFO + ae_per_voy * HARGA_MDO

    # Port charges per voyage
    n_calls = 1 + len(cluster_ports)
    labuh = GT * 83 * n_calls
    pandu = (99_500 + GT * 28) * n_calls
    tunda = compute_biaya_tunda(GT) * n_calls
    tambat_origin = GT * 80 * math.ceil(load_days)
    tambat_dest = sum(GT * 80 * math.ceil(unload_times[p]) for p in cluster_ports)
    tambat = tambat_origin + tambat_dest
    port_charges = labuh + pandu + tunda + tambat

    # TCH annual
    tch_daily_usd = ship['DWT'] * 0.0401 + 18_318
    tch_annual = tch_daily_usd * KURS * COMMISSION_DAYS  # on-hire days only

    # Not used by the model or by any reported result (kept only so that the
    # pre-computation record is unchanged from the production runs).
    total_annual_cost = tch_annual + fuel_per_voy * max_freq + port_charges * max_freq
    max_annual_cap = payload * max_freq
    oc_rate = total_annual_cost / max_annual_cap if max_annual_cap > 0 else 0

    return {
        'feasible': True,
        'D_c': D_c,
        'arrival_offsets': arrival_offsets,
        'RTD_periods': RTD_periods,
        'RTD_days': RTD_days,
        'max_freq': max_freq,
        'total_sail_dist': total_sail_dist,
        'load_days': load_days,
        'unload_times': unload_times,
        'sea_days': sea_days,
        'port_days': port_days,
        'fuel_per_voy': fuel_per_voy,
        'port_charges': port_charges,
        'labuh': labuh, 'pandu': pandu, 'tunda': tunda, 'tambat': tambat,
        'tch_annual': tch_annual,
        'oc_rate': oc_rate,
    }


# =============================================================
# STAGE 2: INTEGRATED MILP
# =============================================================

def build_and_solve(clusters, cluster_routes, precomp):
    """
    Build and solve the integrated time-space MILP.

    Variables:
      y[v]       ∈ {0,1}  : vessel v is used (TCH counted once)
      a[v,c]     ∈ {0,1}  : vessel v assigned to cluster c (multi-cluster OK)
      z[v,c,t]   ∈ {0,1}  : vessel v departs on cluster c route at period t
      q[v,c,p,t] ≥ 0      : cargo (tons) for port p, vessel v departing at t
      s[p,t]     ≥ 0      : inventory at port p at end of period t
      so[p,t]    ≥ 0      : stockout (unmet demand) at port p in period t

    Key: vessels can serve MULTIPLE clusters; cross-cluster non-overlap enforced.

    Objective: min TCH + VoyageCost + CHC + HoldingCost + StockoutPenalty  (Eq. 3)
    """
    n_clusters = len(clusters)

    # ── Demand per period ──
    demand_per_period = {}
    for p_name in dest_port_names:
        demand_per_period[p_name] = ports[p_name]['demand'] * PERIOD_DAYS / TOTAL_DAYS

    # ── Identify feasible (vessel, cluster) pairs ──
    feasible = {}  # (v, c) → precomp data
    for v in range(N_SHIPS):
        for c in range(n_clusters):
            pc = precomp[c].get(v)
            if pc and pc['feasible']:
                feasible[(v, c)] = pc

    print(f"\n  Feasible (vessel, cluster) pairs: {len(feasible)}")
    for c in range(n_clusters):
        fv = [ships[v]['name'] for v in range(N_SHIPS) if (v, c) in feasible]
        print(f"    Cluster {c+1}: {len(fv)} vessels: {', '.join(fv)}")

    # ── Max departure period per (v,c) ──
    max_dep = {}
    for (v, c), pc in feasible.items():
        max_dep[(v, c)] = T_MAX - pc['RTD_periods']

    # ── Port-to-cluster mapping ──
    port_cluster = {}
    for c, cl_ports in enumerate(clusters):
        for p in cl_ports:
            port_cluster[p] = c

    # ================================================================
    # BUILD MODEL
    # ================================================================

    if USE_GUROBI:
        return _build_gurobi(clusters, cluster_routes, precomp, feasible,
                             max_dep, port_cluster, demand_per_period)
    else:
        return _build_ortools(clusters, cluster_routes, precomp, feasible,
                              max_dep, port_cluster, demand_per_period)


def _build_gurobi(clusters, cluster_routes, precomp, feasible,
                   max_dep, port_cluster, demand_per_period):
    """Build and solve with Gurobi."""
    n_clusters = len(clusters)

    model = gp.Model("MIRP_OptionC")
    model.Params.TimeLimit = TIME_LIMIT_SEC
    model.Params.MIPGap = 0.02      # 2% optimality gap acceptable
    model.Params.Threads = 0        # Use all available threads

    # ── Variables ──

    # y[v]: vessel v is used (binary); TCH counted once per vessel
    y = {}
    for v in range(N_SHIPS):
        y[v] = model.addVar(vtype=GRB.BINARY, name=f'y_{v}')

    # a[v,c]: vessel v assigned to cluster c (can serve MULTIPLE clusters)
    a = {}
    for v in range(N_SHIPS):
        for c in range(n_clusters):
            a[v, c] = model.addVar(vtype=GRB.BINARY, name=f'a_{v}_{c}')

    # z[v,c,t]: vessel v departs on cluster c at period t
    z = {}
    for (v, c), pc in feasible.items():
        for t in range(max_dep[(v, c)] + 1):
            z[v, c, t] = model.addVar(vtype=GRB.BINARY, name=f'z_{v}_{c}_{t}')

    # q[v,c,p,t]: cargo for port p when vessel v departs cluster c at period t
    q = {}
    for (v, c), pc in feasible.items():
        for p_name in clusters[c]:
            for t in range(max_dep[(v, c)] + 1):
                arr_t = t + pc['arrival_offsets'][p_name]
                if arr_t < T_MAX:
                    q[v, c, p_name, t] = model.addVar(
                        lb=0, ub=ships[v]['payload'],
                        name=f'q_{v}_{c}_{p_name}_{t}')

    # s[p,t]: inventory at port p at end of period t
    s = {}
    for p_name in dest_port_names:
        for t in range(T_MAX):
            s[p_name, t] = model.addVar(
                lb=0, ub=ports[p_name]['capacity'],
                name=f's_{p_name}_{t}')

    # so[p,t]: stockout at port p in period t
    so = {}
    for p_name in dest_port_names:
        for t in range(T_MAX):
            so[p_name, t] = model.addVar(lb=0, name=f'so_{p_name}_{t}')

    model.update()

    print(f"\n  Variables created:")
    print(f"    y (vessel used) : {len(y)}")
    print(f"    a (assignment)  : {len(a)}")
    print(f"    z (departures)  : {len(z)}")
    print(f"    q (cargo alloc) : {len(q)}")
    print(f"    s (inventory)   : {len(s)}")
    print(f"    so (stockout)   : {len(so)}")
    print(f"    Total           : {len(y)+len(a)+len(z)+len(q)+len(s)+len(so)}")

    # ── Constraints ──

    # (C1) Link y[v] to a[v,c]: y[v] >= a[v,c]; vessel is "used" if assigned
    #      to any cluster. (No upper limit on how many clusters a vessel serves.)
    for v in range(N_SHIPS):
        for c in range(n_clusters):
            model.addConstr(y[v] >= a[v, c], name=f'link_y_{v}_{c}')

    # (C2) Infeasible pairs blocked
    for v in range(N_SHIPS):
        for c in range(n_clusters):
            if (v, c) not in feasible:
                model.addConstr(a[v, c] == 0, name=f'infeas_{v}_{c}')

    # (C3) Each cluster served by at least 1 vessel
    for c in range(n_clusters):
        model.addConstr(
            gp.quicksum(a[v, c] for v in range(N_SHIPS)) >= 1,
            name=f'serve_{c}')

    # (C4) Departure only if assigned: z[v,c,t] <= a[v,c]
    for (v, c, t) in z:
        model.addConstr(z[v, c, t] <= a[v, c], name=f'link_z_{v}_{c}_{t}')

    # (C5) Cross-cluster non-overlap: at each period τ, at most 1 active
    #      voyage per vessel across ALL clusters.
    #      A departure z[v,c,t] occupies vessel v for periods t..t+RTD[v,c]-1.
    #      So at period τ: sum over all (c, t) where t <= τ < t+RTD[v,c] of z[v,c,t] <= 1
    for v in range(N_SHIPS):
        # Collect all feasible (c, max_dep) for this vessel
        v_clusters = [(c, pc) for (vv, c), pc in feasible.items() if vv == v]
        if not v_clusters:
            continue
        for tau in range(T_MAX):
            active_vars = []
            for c, pc in v_clusters:
                rtd = pc['RTD_periods']
                # Departure at time t occupies [t, t+rtd).
                # Active at τ means t <= τ < t + rtd, i.e. t in [τ-rtd+1, τ]
                t_lo = max(0, tau - rtd + 1)
                t_hi = min(tau, max_dep.get((v, c), -1))
                for t in range(t_lo, t_hi + 1):
                    if (v, c, t) in z:
                        active_vars.append(z[v, c, t])
            if len(active_vars) > 1:
                model.addConstr(
                    gp.quicksum(active_vars) <= 1,
                    name=f'noovlp_{v}_{tau}')

    # (C6) Cargo capacity per voyage: sum_p q[v,c,p,t] <= K_v * z[v,c,t]
    for (v, c), pc in feasible.items():
        for t in range(max_dep[(v, c)] + 1):
            cargo_vars = [q[v, c, p, t] for p in clusters[c]
                         if (v, c, p, t) in q]
            if cargo_vars and (v, c, t) in z:
                model.addConstr(
                    gp.quicksum(cargo_vars) <= ships[v]['payload'] * z[v, c, t],
                    name=f'cap_{v}_{c}_{t}')

    # (C7) Inventory balance: s[p,t] = s[p,t-1] + deliveries - demand + stockout
    for p_name in dest_port_names:
        c = port_cluster[p_name]
        for t in range(T_MAX):
            # Deliveries arriving at port p in period t
            delivery_expr = gp.LinExpr()
            for v in range(N_SHIPS):
                if (v, c) not in feasible:
                    continue
                pc = feasible[(v, c)]
                arr_offset = pc['arrival_offsets'][p_name]
                dep_t = t - arr_offset  # departure period that results in arrival at t
                if dep_t >= 0 and (v, c, p_name, dep_t) in q:
                    delivery_expr.add(q[v, c, p_name, dep_t])

            if t == 0:
                # Initial inventory
                model.addConstr(
                    s[p_name, 0] == ports[p_name]['initial_stock']
                    + delivery_expr
                    - demand_per_period[p_name]
                    + so[p_name, 0],
                    name=f'inv_{p_name}_{t}')
            else:
                model.addConstr(
                    s[p_name, t] == s[p_name, t-1]
                    + delivery_expr
                    - demand_per_period[p_name]
                    + so[p_name, t],
                    name=f'inv_{p_name}_{t}')

    # ── Objective Function ──
    obj = gp.LinExpr()

    # (1) TCH: annual charter cost, counted ONCE per vessel used (y[v])
    for v in range(N_SHIPS):
        tch_v = (ships[v]['DWT'] * 0.0401 + 18_318) * KURS * COMMISSION_DAYS  # on-hire days only
        obj.add(y[v], tch_v)

    # (2) Voyage cost: fuel + port charges per departure
    for (v, c, t) in z:
        pc = feasible[(v, c)]
        voy_cost = pc['fuel_per_voy'] + pc['port_charges']
        obj.add(z[v, c, t], voy_cost)

    # (3) CHC: cargo handling cost per ton delivered
    for (v, c, p, t) in q:
        obj.add(q[v, c, p, t], CHC_RATE)

    # (4) Holding cost: per ton per period
    for p_name in dest_port_names:
        hc_rate = ports[p_name]['holding_cost'] * PERIOD_DAYS
        for t in range(T_MAX):
            obj.add(s[p_name, t], hc_rate)

    # (5) Stockout penalty
    for p_name in dest_port_names:
        for t in range(T_MAX):
            obj.add(so[p_name, t], STOCKOUT_PENALTY)

    # The objective contains no capacity- or utilization-based penalty term;
    # cargo-space utilization (Eq. 8) is computed after solving, for reporting.

    model.setObjective(obj, GRB.MINIMIZE)

    # ── Solve ──
    print(f"\n  Solving with Gurobi (time limit: {TIME_LIMIT_SEC}s)...")
    print(f"  Total variables   : {model.NumVars}")
    print(f"  Binary variables  : {model.NumIntVars}")
    print(f"  Constraints       : {model.NumConstrs}")

    _wall_start = _time.time()
    model.optimize()
    _wall_elapsed = _time.time() - _wall_start
    _gurobi_runtime = model.Runtime

    print(f"  Wall-clock time : {_wall_elapsed:,.1f} s")
    print(f"  Gurobi Runtime  : {_gurobi_runtime:,.1f} s")

    if model.Status == GRB.INFEASIBLE:
        print("  MODEL INFEASIBLE")
        model.computeIIS()
        model.write("infeasible.ilp")
        return None

    if model.SolCount == 0:
        print("  NO SOLUTION FOUND")
        return None

    status_str = {
        GRB.OPTIMAL: "OPTIMAL",
        GRB.TIME_LIMIT: f"TIME_LIMIT (gap={model.MIPGap:.1%})",
    }.get(model.Status, f"Status={model.Status}")
    print(f"  Status: {status_str}")
    print(f"  Objective: Rp {model.ObjVal:,.0f}")

    # ── Extract Solution ──
    sol = _extract_solution_gurobi(model, y, a, z, q, s, so,
                                    clusters, cluster_routes, precomp,
                                    feasible, demand_per_period)
    if sol is not None:
        sol['wall_time_sec'] = _wall_elapsed
        sol['gurobi_runtime_sec'] = _gurobi_runtime
        # Publication-grade solver metadata (taken directly from the model)
        sol['status'] = status_str
        sol['obj_bound'] = model.ObjBound      # best bound (NOT reconstructed from gap)
        sol['sol_count'] = model.SolCount
    return sol


def _extract_solution_gurobi(model, y, a, z, q, s, so,
                              clusters, cluster_routes, precomp,
                              feasible, demand_per_period):
    """Extract and format solution from Gurobi model."""
    n_clusters = len(clusters)
    sol = {
        'objective': model.ObjVal,
        'gap': model.MIPGap if hasattr(model, 'MIPGap') else 0,
        'assignments': [],
        'schedules': {},    # v → list of departure periods
        'cargo': {},        # (v, dep_t) → {port: tons}
        'inventory': {},    # port → [s_0, s_1, ..., s_T]
        'stockout': {},     # port → [so_0, so_1, ..., so_T]
        'costs': {},
    }

    # Assignments
    for v in range(N_SHIPS):
        for c in range(n_clusters):
            if a[v, c].X > 0.5:
                sol['assignments'].append((v, c))

    # Schedules
    for v in range(N_SHIPS):
        sol['schedules'][v] = []
    for (v, c, t), var in z.items():
        if var.X > 0.5:
            sol['schedules'][v].append((c, t))
    for v in sol['schedules']:
        sol['schedules'][v].sort(key=lambda x: x[1])

    # Cargo allocation
    for (v, c, p, t), var in q.items():
        if var.X > 0.1:
            key = (v, c, t)
            if key not in sol['cargo']:
                sol['cargo'][key] = {}
            sol['cargo'][key][p] = var.X

    # Inventory
    for p_name in dest_port_names:
        sol['inventory'][p_name] = [s[p_name, t].X for t in range(T_MAX)]
        sol['stockout'][p_name] = [so[p_name, t].X for t in range(T_MAX)]

    # Cost breakdown
    total_tch = 0
    total_fuel = 0
    total_port = 0
    total_chc = 0
    total_hc = 0
    total_so = 0

    # TCH: count once per vessel used (y[v])
    for v in range(N_SHIPS):
        if y[v].X > 0.5:
            tch_v = (ships[v]['DWT'] * 0.0401 + 18_318) * KURS * COMMISSION_DAYS  # on-hire days only
            total_tch += tch_v

    cap_deployed_tc = 0.0   # SUM_{v,c,t} K_v z_{v,c,t} over deployed TC voyages
    for (v, c, t), var in z.items():
        if var.X > 0.5:
            pc = feasible[(v, c)]
            total_fuel += pc['fuel_per_voy']
            total_port += pc['port_charges']
            cap_deployed_tc += ships[v]['payload']

    q_delivered_tc = 0.0    # SUM_{v,c,p,t} q_{v,c,p,t}
    for (v, c, p, t), var in q.items():
        if var.X > 0.1:
            total_chc += CHC_RATE * var.X
            q_delivered_tc += var.X

    for p_name in dest_port_names:
        hc_rate = ports[p_name]['holding_cost'] * PERIOD_DAYS
        for t in range(T_MAX):
            total_hc += hc_rate * s[p_name, t].X
            total_so += STOCKOUT_PENALTY * so[p_name, t].X

    financial_stage2 = total_tch + total_fuel + total_port + total_chc + total_hc
    sol['costs'] = {
        'tch': total_tch,
        'fuel': total_fuel,
        'port_charges': total_port,
        'chc': total_chc,
        'holding': total_hc,
        'stockout': total_so,
        'financial_stage2': financial_stage2,   # direct expenditure, excludes penalties
        'total': financial_stage2 + total_so,    # = decision objective (financial + penalty)
    }
    # Cargo-space utilization KPI (deployed TC voyages only), Eq. (8): U = SUM q / SUM (K*z)
    sol['kpi'] = {
        'Q_TC': q_delivered_tc,
        'Cap_deployed_TC': cap_deployed_tc,
        'U_cargo_TC': (q_delivered_tc / cap_deployed_tc) if cap_deployed_tc > 0 else 0.0,
    }

    return sol


def _build_ortools(clusters, cluster_routes, precomp, feasible,
                    max_dep, port_cluster, demand_per_period):
    """Build and solve with OR-Tools (fallback)."""
    n_clusters = len(clusters)

    solver = pywraplp.Solver.CreateSolver('SCIP')
    if not solver:
        solver = pywraplp.Solver.CreateSolver('CBC')
    if not solver:
        print("ERROR: No solver available")
        return None

    solver.SetTimeLimit(TIME_LIMIT_SEC * 1000)

    # ── Variables ──

    # y[v]: vessel v is used (binary); TCH counted once
    y = {}
    for v in range(N_SHIPS):
        y[v] = solver.BoolVar(f'y_{v}')

    # a[v,c]: vessel v assigned to cluster c (can serve MULTIPLE clusters)
    a = {}
    for v in range(N_SHIPS):
        for c in range(n_clusters):
            a[v, c] = solver.BoolVar(f'a_{v}_{c}')

    z = {}
    for (v, c), pc in feasible.items():
        for t in range(max_dep[(v, c)] + 1):
            z[v, c, t] = solver.BoolVar(f'z_{v}_{c}_{t}')

    q = {}
    for (v, c), pc in feasible.items():
        for p_name in clusters[c]:
            for t in range(max_dep[(v, c)] + 1):
                arr_t = t + pc['arrival_offsets'][p_name]
                if arr_t < T_MAX:
                    q[v, c, p_name, t] = solver.NumVar(
                        0, ships[v]['payload'], f'q_{v}_{c}_{p_name}_{t}')

    s = {}
    for p_name in dest_port_names:
        for t in range(T_MAX):
            s[p_name, t] = solver.NumVar(
                0, ports[p_name]['capacity'], f's_{p_name}_{t}')

    so = {}
    for p_name in dest_port_names:
        for t in range(T_MAX):
            so[p_name, t] = solver.NumVar(0, solver.infinity(), f'so_{p_name}_{t}')

    print(f"\n  Variables: {solver.NumVariables()}")

    # ── Constraints ──

    # (C1) Link y[v] >= a[v,c]: vessel "used" if assigned to any cluster
    for v in range(N_SHIPS):
        for c in range(n_clusters):
            solver.Add(y[v] >= a[v, c])

    # (C2) Infeasible pairs
    for v in range(N_SHIPS):
        for c in range(n_clusters):
            if (v, c) not in feasible:
                solver.Add(a[v, c] == 0)

    # (C3) Each cluster served
    for c in range(n_clusters):
        solver.Add(sum(a[v, c] for v in range(N_SHIPS)) >= 1)

    # (C4) z <= a
    for (v, c, t) in z:
        solver.Add(z[v, c, t] <= a[v, c])

    # (C5) Cross-cluster non-overlap: at each period τ, at most 1 active voyage
    for v in range(N_SHIPS):
        v_clusters = [(c, pc) for (vv, c), pc in feasible.items() if vv == v]
        if not v_clusters:
            continue
        for tau in range(T_MAX):
            active_vars = []
            for c, pc in v_clusters:
                rtd = pc['RTD_periods']
                t_lo = max(0, tau - rtd + 1)
                t_hi = min(tau, max_dep.get((v, c), -1))
                for t in range(t_lo, t_hi + 1):
                    if (v, c, t) in z:
                        active_vars.append(z[v, c, t])
            if len(active_vars) > 1:
                solver.Add(sum(active_vars) <= 1)

    # (C6) Cargo capacity
    for (v, c), pc in feasible.items():
        for t in range(max_dep[(v, c)] + 1):
            cargo = [q[v, c, p, t] for p in clusters[c] if (v, c, p, t) in q]
            if cargo and (v, c, t) in z:
                solver.Add(sum(cargo) <= ships[v]['payload'] * z[v, c, t])

    # (C7) Inventory balance
    for p_name in dest_port_names:
        c = port_cluster[p_name]
        for t in range(T_MAX):
            delivery = []
            for v in range(N_SHIPS):
                if (v, c) not in feasible:
                    continue
                pc = feasible[(v, c)]
                arr_off = pc['arrival_offsets'][p_name]
                dep_t = t - arr_off
                if dep_t >= 0 and (v, c, p_name, dep_t) in q:
                    delivery.append(q[v, c, p_name, dep_t])

            if t == 0:
                solver.Add(
                    s[p_name, 0] == ports[p_name]['initial_stock']
                    + sum(delivery)
                    - demand_per_period[p_name]
                    + so[p_name, 0])
            else:
                solver.Add(
                    s[p_name, t] == s[p_name, t - 1]
                    + sum(delivery)
                    - demand_per_period[p_name]
                    + so[p_name, t])

    # ── Objective ──
    obj = solver.Objective()

    # TCH counted once per vessel used (y[v])
    for v in range(N_SHIPS):
        tch_v = (ships[v]['DWT'] * 0.0401 + 18_318) * KURS * COMMISSION_DAYS  # on-hire days only
        obj.SetCoefficient(y[v], tch_v)

    for (v, c, t) in z:
        pc = feasible[(v, c)]
        voy_cost = pc['fuel_per_voy'] + pc['port_charges']
        obj.SetCoefficient(z[v, c, t], voy_cost)

    for (v, c, p, t) in q:
        obj.SetCoefficient(q[v, c, p, t], CHC_RATE)

    for p_name in dest_port_names:
        hc_rate = ports[p_name]['holding_cost'] * PERIOD_DAYS
        for t in range(T_MAX):
            obj.SetCoefficient(s[p_name, t], hc_rate)
            obj.SetCoefficient(so[p_name, t], STOCKOUT_PENALTY)

    obj.SetMinimization()

    print(f"  Constraints: {solver.NumConstraints()}")
    print(f"  Solving with OR-Tools (time limit: {TIME_LIMIT_SEC}s)...")

    status = solver.Solve()
    if status not in (pywraplp.Solver.OPTIMAL, pywraplp.Solver.FEASIBLE):
        print(f"  NO SOLUTION (status={status})")
        return None

    status_str = 'OPTIMAL' if status == pywraplp.Solver.OPTIMAL else 'FEASIBLE'
    print(f"  Status: {status_str}")
    print(f"  Objective: Rp {solver.Objective().Value():,.0f}")

    # ── Extract Solution ──
    return _extract_solution_ortools(solver, y, a, z, q, s, so,
                                      clusters, cluster_routes, precomp,
                                      feasible, demand_per_period)


def _extract_solution_ortools(solver, y, a, z, q, s, so,
                               clusters, cluster_routes, precomp,
                               feasible, demand_per_period):
    """Extract solution from OR-Tools."""
    n_clusters = len(clusters)
    sol = {
        'objective': solver.Objective().Value(),
        'gap': 0,
        'assignments': [],
        'schedules': {},
        'cargo': {},
        'inventory': {},
        'stockout': {},
        'costs': {},
    }

    for v in range(N_SHIPS):
        for c in range(n_clusters):
            if a[v, c].solution_value() > 0.5:
                sol['assignments'].append((v, c))

    for v in range(N_SHIPS):
        sol['schedules'][v] = []
    for (v, c, t), var in z.items():
        if var.solution_value() > 0.5:
            sol['schedules'][v].append((c, t))
    for v in sol['schedules']:
        sol['schedules'][v].sort(key=lambda x: x[1])

    for (v, c, p, t), var in q.items():
        if var.solution_value() > 0.1:
            key = (v, c, t)
            if key not in sol['cargo']:
                sol['cargo'][key] = {}
            sol['cargo'][key][p] = var.solution_value()

    for p_name in dest_port_names:
        sol['inventory'][p_name] = [s[p_name, t].solution_value() for t in range(T_MAX)]
        sol['stockout'][p_name] = [so[p_name, t].solution_value() for t in range(T_MAX)]

    # Cost breakdown
    total_tch = 0
    total_fuel = 0
    total_port = 0
    total_chc = 0
    total_hc = 0
    total_so_cost = 0

    # TCH: count once per vessel used (y[v])
    for v in range(N_SHIPS):
        if y[v].solution_value() > 0.5:
            tch_v = (ships[v]['DWT'] * 0.0401 + 18_318) * KURS * COMMISSION_DAYS  # on-hire days only
            total_tch += tch_v

    cap_deployed_tc = 0.0
    for (v, c, t), var in z.items():
        if var.solution_value() > 0.5:
            pc = feasible[(v, c)]
            total_fuel += pc['fuel_per_voy']
            total_port += pc['port_charges']
            cap_deployed_tc += ships[v]['payload']

    q_delivered_tc = 0.0
    for (v, c, p, t), var in q.items():
        if var.solution_value() > 0.1:
            total_chc += CHC_RATE * var.solution_value()
            q_delivered_tc += var.solution_value()

    for p_name in dest_port_names:
        hc_rate = ports[p_name]['holding_cost'] * PERIOD_DAYS
        for t in range(T_MAX):
            total_hc += hc_rate * s[p_name, t].solution_value()
            total_so_cost += STOCKOUT_PENALTY * so[p_name, t].solution_value()

    financial_stage2 = total_tch + total_fuel + total_port + total_chc + total_hc
    sol['costs'] = {
        'tch': total_tch, 'fuel': total_fuel, 'port_charges': total_port,
        'chc': total_chc, 'holding': total_hc, 'stockout': total_so_cost,
        'financial_stage2': financial_stage2,
        'total': financial_stage2 + total_so_cost,
    }
    sol['kpi'] = {
        'Q_TC': q_delivered_tc,
        'Cap_deployed_TC': cap_deployed_tc,
        'U_cargo_TC': (q_delivered_tc / cap_deployed_tc) if cap_deployed_tc > 0 else 0.0,
    }

    return sol


# =============================================================
# OUTPUT GENERATION
# =============================================================

def generate_itinerary_xlsx(sol, clusters, cluster_routes, precomp, feasible, out_dir):
    """Generate Excel file with per-vessel voyage itinerary (one sheet per vessel)."""
    if not HAS_OPENPYXL:
        print("  [SKIP] openpyxl not installed; itinerary Excel not generated")
        return

    wb = Workbook()
    wb.remove(wb.active)

    header_font = Font(name='Arial', bold=True, size=11)
    header_fill = PatternFill('solid', fgColor='4472C4')
    header_font_white = Font(name='Arial', bold=True, size=10, color='FFFFFF')
    data_font = Font(name='Arial', size=10)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'))

    for v in range(N_SHIPS):
        voyages = sol['schedules'][v]
        if not voyages:
            continue

        ship = ships[v]
        sheet_name = f"{ship['id']} {ship['name']}"[:31]
        ws = wb.create_sheet(title=sheet_name)

        total_cargo_all = 0
        for (c, t) in voyages:
            total_cargo_all += sum(sol['cargo'].get((v, c, t), {}).values())

        ws['A1'] = f"{ship['id']} ({ship['name']})"
        ws['A1'].font = Font(name='Arial', bold=True, size=12)
        ws['B1'] = 'No. of Voyages:'
        ws['B1'].font = header_font
        ws['C1'] = len(voyages)
        ws['C1'].font = header_font
        ws['D1'] = 'Total cargo carried:'
        ws['D1'].font = header_font
        ws['E1'] = round(total_cargo_all, 0)
        ws['E1'].font = header_font
        ws['E1'].number_format = '#,##0'
        ws['F1'] = 'ton'
        ws['F1'].font = header_font

        headers = ['Voy. No.', 'Cluster', 'Day', 'Origin', 'Destination', 'Cargo Delivered (ton)']
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=2, column=col_idx, value=h)
            cell.font = header_font_white
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border

        row_num = 3
        for voy_idx, (c, t) in enumerate(voyages, 1):
            route = cluster_routes[c]
            cargo = sol['cargo'].get((v, c, t), {})
            pc = feasible.get((v, c))
            dep_day = t * PERIOD_DAYS

            ws.cell(row=row_num, column=1, value=voy_idx).font = data_font
            ws.cell(row=row_num, column=2, value=c + 1).font = data_font
            ws.cell(row=row_num, column=3, value=dep_day).font = data_font
            ws.cell(row=row_num, column=4, value='Biringkassi').font = data_font
            ws.cell(row=row_num, column=5, value=route[0]).font = data_font
            cargo_val = round(cargo.get(route[0], 0), 1)
            cell_f = ws.cell(row=row_num, column=6, value=cargo_val)
            cell_f.font = data_font
            cell_f.number_format = '#,##0.0'
            for col_i in range(1, 7):
                ws.cell(row=row_num, column=col_i).border = thin_border
            row_num += 1

            for leg_idx in range(1, len(route)):
                prev_port = route[leg_idx - 1]
                next_port = route[leg_idx]
                if pc:
                    arr_day = (t + pc['arrival_offsets'].get(next_port, 0)) * PERIOD_DAYS
                else:
                    arr_day = dep_day

                ws.cell(row=row_num, column=1, value=voy_idx).font = data_font
                ws.cell(row=row_num, column=2, value=c + 1).font = data_font
                ws.cell(row=row_num, column=3, value=arr_day).font = data_font
                ws.cell(row=row_num, column=4, value=prev_port).font = data_font
                ws.cell(row=row_num, column=5, value=next_port).font = data_font
                cargo_val = round(cargo.get(next_port, 0), 1)
                cell_f = ws.cell(row=row_num, column=6, value=cargo_val)
                cell_f.font = data_font
                cell_f.number_format = '#,##0.0'
                for col_i in range(1, 7):
                    ws.cell(row=row_num, column=col_i).border = thin_border
                row_num += 1

            last_port = route[-1]
            if pc:
                rtd_days = pc['RTD_days']
                return_day = round(dep_day + rtd_days)
            else:
                return_day = dep_day
            ws.cell(row=row_num, column=1, value=voy_idx).font = data_font
            ws.cell(row=row_num, column=2, value=c + 1).font = data_font
            ws.cell(row=row_num, column=3, value=return_day).font = data_font
            ws.cell(row=row_num, column=4, value=last_port).font = data_font
            ws.cell(row=row_num, column=5, value='Biringkassi').font = data_font
            ws.cell(row=row_num, column=6, value=0).font = data_font
            ws.cell(row=row_num, column=6).number_format = '#,##0.0'
            for col_i in range(1, 7):
                ws.cell(row=row_num, column=col_i).border = thin_border
            row_num += 1

        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 9
        ws.column_dimensions['C'].width = 8
        ws.column_dimensions['D'].width = 18
        ws.column_dimensions['E'].width = 18
        ws.column_dimensions['F'].width = 22

    xlsx_path = os.path.join(out_dir, f'OptionC_k{K_CLUSTERS}_Itinerary.xlsx')
    wb.save(xlsx_path)
    print(f"  Itinerary Excel saved: {xlsx_path}")


def generate_report(sol, clusters, cluster_routes, precomp, feasible, out_dir):
    """Generate text report, CSVs, and inventory charts."""
    os.makedirs(out_dir, exist_ok=True)

    costs = sol['costs']
    rpt = []
    rpt.append("=" * 90)
    rpt.append("MIRP OPTION C: Integrated Scheduling, Fleet Sizing & Inventory Control")
    rpt.append("=" * 90)
    rpt.append(f"Model       : Time-indexed MILP with inventory constraints (multi-cluster)")
    rpt.append(f"Solver      : {'Gurobi' if USE_GUROBI else 'OR-Tools'}")
    rpt.append(f"Time periods: {T_MAX} ({PERIOD_DAYS}-day periods)")
    rpt.append(f"Fleet pool  : {N_SHIPS} existing vessels (CC1-CC8)")
    rpt.append(f"Clusters    : {len(clusters)}")
    rpt.append(f"KURS        : Rp {KURS:,}/USD  (final study rate)")
    rpt.append(f"Solver Status : {sol.get('status', 'NA')}")
    rpt.append(f"Objective (incumbent ObjVal): Rp {sol['objective']:,.0f}")
    if sol.get('obj_bound') is not None:
        rpt.append(f"Best Bound (ObjBound)       : Rp {sol['obj_bound']:,.0f}")
    if sol.get('gap') is not None:
        rpt.append(f"MIP Gap     : {sol['gap']:.4%}")
    if sol.get('sol_count') is not None:
        rpt.append(f"Solution Count: {sol['sol_count']}")
    if sol.get('gurobi_runtime_sec') is not None:
        rpt.append(f"Gurobi Runtime: {sol['gurobi_runtime_sec']:,.1f} seconds")
    if sol.get('wall_time_sec') is not None:
        rpt.append(f"Wall-clock    : {sol['wall_time_sec']:,.1f} seconds")

    # Clusters
    rpt.append(f"\nCLUSTERS & ROUTES:")
    for c, cl_ports in enumerate(clusters):
        D_c = sum(ports[p]['demand'] for p in cl_ports)
        route_str = " -> ".join(["Biringkassi"] + cluster_routes[c] + ["Biringkassi"])
        rpt.append(f"  Cluster {c+1}: {', '.join(cl_ports)}")
        rpt.append(f"    Demand: {D_c:,} ton/year | Route: {route_str}")

    # Vessel assignments (multi-cluster possible)
    rpt.append(f"\nVESSEL ASSIGNMENTS (multi-cluster enabled):")
    rpt.append(f"  {'Vessel':<22} {'Clusters':>10} {'Payload':>8} {'Tot.Voy':>8}")
    rpt.append("  " + "-" * 55)

    used_vessels = set()
    for v in range(N_SHIPS):
        assigned_clusters = [c for (vv, c) in sol['assignments'] if vv == v]
        n_voy = len(sol['schedules'][v])
        if assigned_clusters:
            used_vessels.add(v)
            cl_str = ','.join(str(c+1) for c in sorted(assigned_clusters))
            # Voyage count per cluster
            voy_per_cl = {}
            for (c, t) in sol['schedules'][v]:
                voy_per_cl[c] = voy_per_cl.get(c, 0) + 1
            detail = ' | '.join(f"Cl{c+1}:{n}voy" for c, n in sorted(voy_per_cl.items()))
            rpt.append(f"  {ships[v]['name']:<22} {cl_str:>10} {ships[v]['payload']:>8,} "
                        f"{n_voy:>8}")
            rpt.append(f"    {detail}")

    rpt.append(f"\n  Vessels used: {len(used_vessels)}/{N_SHIPS}")
    unassigned = [v for v in range(N_SHIPS) if v not in used_vessels]
    if unassigned:
        rpt.append(f"  Unassigned vessels: {', '.join(ships[v]['name'] for v in unassigned)}")

    # Voyage schedule
    rpt.append(f"\nVOYAGE SCHEDULE:")
    rpt.append(f"  {'Vessel':<22} {'Cl':>3} {'Dep.Period':>10} {'Dep.Day':>10} {'Cargo(t)':>10}")
    rpt.append("  " + "-" * 60)

    for v in range(N_SHIPS):
        for (c, t) in sol['schedules'][v]:
            day = t * PERIOD_DAYS
            cargo_key = (v, c, t)
            total_cargo = sum(sol['cargo'].get(cargo_key, {}).values())
            rpt.append(f"  {ships[v]['name']:<22} {c+1:>3} {t:>10} {day:>10} "
                        f"{total_cargo:>10,.0f}")

    # Cost breakdown
    rpt.append(f"\nCOST BREAKDOWN:")
    rpt.append(f"  {'Component':<30} {'Value (Rp)':>25} {'(Million)':>15}")
    rpt.append("  " + "-" * 70)
    for name, key in [('TCH (charter)', 'tch'),
                      ('Fuel Cost', 'fuel'),
                      ('Port Charges', 'port_charges'),
                      ('Cargo Handling (CHC)', 'chc'),
                      ('Holding Cost', 'holding'),
                      ('Stockout Penalty', 'stockout')]:
        val = costs[key]
        rpt.append(f"  {name:<30} {val:>25,.0f} {val/1e6:>15,.1f}")
    rpt.append("  " + "-" * 70)
    rpt.append(f"  {'TOTAL:::::::::::::::::::::::::':<30} {costs['total']:>25,.0f} {costs['total']/1e6:>15,.1f}")
    rpt.append(f"  {'Total (Billion Rp)':<30} {'':>25} {costs['total']/1e9:>15,.3f}")
    rpt.append(f"  {'Financial (direct, excl. penalty)':<30} {costs['financial_stage2']:>25,.0f} {costs['financial_stage2']/1e6:>15,.1f}")

    # Cargo-space utilization KPI (deployed TC voyages only)
    kpi = sol['kpi']
    rpt.append(f"\nCARGO-SPACE UTILIZATION (deployed TC voyages only):")
    rpt.append(f"  Q_TC (delivered)          : {kpi['Q_TC']:>15,.1f} ton")
    rpt.append(f"  Cap_deployed_TC (SUM K*z) : {kpi['Cap_deployed_TC']:>15,.1f} ton")
    rpt.append(f"  U_cargo_TC                : {kpi['U_cargo_TC']*100:>14.2f} %")

    # Inventory summary
    rpt.append(f"\nINVENTORY SUMMARY:")
    rpt.append(f"  {'Port':<20} {'Max Stock':>10} {'Min Stock':>10} "
               f"{'Total SO(t)':>12} {'SO Cost(Rp)':>20}")
    rpt.append("  " + "-" * 75)

    for p_name in dest_port_names:
        inv = sol['inventory'][p_name]
        so_vals = sol['stockout'][p_name]
        max_s = max(inv)
        min_s = min(inv)
        total_so = sum(so_vals)
        so_cost = total_so * STOCKOUT_PENALTY
        rpt.append(f"  {p_name:<20} {max_s:>10,.0f} {min_s:>10,.0f} "
                    f"{total_so:>12,.1f} {so_cost:>20,.0f}")

    report_text = "\n".join(rpt)
    print(report_text)

    with open(os.path.join(out_dir, f'OptionC_k{K_CLUSTERS}_Report.txt'), 'w') as f:
        f.write(report_text)

    # ── Inventory Charts ──
    fig, axes = plt.subplots(3, 3, figsize=(18, 14))
    fig.suptitle(f'Inventory Tracking: Option C Integrated MILP '
                 f'({PERIOD_DAYS}-day periods, T={T_MAX})',
                 fontsize=13, fontweight='bold')

    for idx, p_name in enumerate(dest_port_names):
        ax = axes[idx // 3][idx % 3]
        inv = sol['inventory'][p_name]
        so_vals = sol['stockout'][p_name]
        total_so = sum(so_vals)

        periods = list(range(T_MAX))
        days = [t * PERIOD_DAYS for t in periods]

        ax.fill_between(days, inv, alpha=0.25, color='steelblue')
        ax.plot(days, inv, 'b-', linewidth=0.8, label='Stock Level')
        ax.axhline(y=ports[p_name]['capacity'], color='r', linestyle='--',
                   linewidth=0.7, label=f"Cap ({ports[p_name]['capacity']:,})")
        ax.axhline(y=0, color='gray', linewidth=0.5)

        # Mark delivery periods
        c_idx = None
        for (v, c) in sol['assignments']:
            if p_name in clusters[c]:
                c_idx = c
                break
        if c_idx is not None:
            for v in range(N_SHIPS):
                for (c, t) in sol['schedules'][v]:
                    if c == c_idx and (v, c) in feasible:
                        arr_off = feasible[(v, c)]['arrival_offsets'].get(p_name, 0)
                        arr_day = (t + arr_off) * PERIOD_DAYS
                        if arr_day < TOTAL_DAYS:
                            ax.axvline(x=arr_day, color='green', alpha=0.4, linewidth=0.5)

        so_lbl = f"SO: {total_so:.0f}t" if total_so > 0.1 else "No SO"
        ax.set_title(f"{p_name} | Cap={ports[p_name]['capacity']:,}t\n{so_lbl}", fontsize=8)
        ax.set_xlabel('Day', fontsize=7)
        ax.set_ylabel('Ton', fontsize=7)
        ax.tick_params(labelsize=6)
        ax.legend(fontsize=6, loc='upper right')
        ax.set_xlim(0, TOTAL_DAYS)
        ax.grid(True, alpha=0.3)

    plt.tight_layout()
    plt.savefig(os.path.join(out_dir, f'OptionC_k{K_CLUSTERS}_Inventory.png'),
                dpi=150, bbox_inches='tight')
    plt.close()

    # ── Schedule CSV ──
    with open(os.path.join(out_dir, f'OptionC_k{K_CLUSTERS}_Schedule.csv'), 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['Vessel', 'Cluster', 'Dep_Period', 'Dep_Day'] +
                   [f'Cargo_{p}' for p in dest_port_names] + ['Total_Cargo'])
        for v in range(N_SHIPS):
            for (c, t) in sol['schedules'][v]:
                cargo = sol['cargo'].get((v, c, t), {})
                row = [ships[v]['name'], c+1, t, t*PERIOD_DAYS]
                total = 0
                for p in dest_port_names:
                    val = cargo.get(p, 0)
                    row.append(round(val, 1))
                    total += val
                row.append(round(total, 1))
                w.writerow(row)

    # ── Inventory CSV ──
    with open(os.path.join(out_dir, f'OptionC_k{K_CLUSTERS}_Inventory.csv'), 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['Period', 'Day'] + [f's_{p}' for p in dest_port_names] +
                   [f'so_{p}' for p in dest_port_names])
        for t in range(T_MAX):
            row = [t, t * PERIOD_DAYS]
            for p in dest_port_names:
                row.append(round(sol['inventory'][p][t], 1))
            for p in dest_port_names:
                row.append(round(sol['stockout'][p][t], 1))
            w.writerow(row)

    # ── Itinerary Excel ──
    generate_itinerary_xlsx(sol, clusters, cluster_routes, precomp, feasible, out_dir)

    print(f"\nFiles saved in {out_dir}")
    return costs['total']


# =============================================================
# MAIN
# =============================================================

# Stage-1 K-Means memberships of the reported production runs (K = 2..9), in the
# exact list order used by Stage 2 (coordinates, Euclidean distance, seed 42).
RECORDED_KMEANS_CLUSTERS = {2: [['Mamuju', 'Palu', 'Kendari', 'Celukan Bawang', 'Lembar'], ['Bitung', 'Ambon', 'Oba', 'Sorong']], 3: [['Bitung', 'Ambon', 'Oba', 'Sorong'], ['Celukan Bawang', 'Lembar'], ['Mamuju', 'Palu', 'Kendari']], 4: [['Mamuju', 'Palu', 'Kendari'], ['Celukan Bawang', 'Lembar'], ['Ambon', 'Sorong'], ['Bitung', 'Oba']], 5: [['Bitung', 'Oba'], ['Mamuju', 'Palu'], ['Celukan Bawang', 'Lembar'], ['Kendari'], ['Ambon', 'Sorong']], 6: [['Bitung', 'Oba'], ['Celukan Bawang', 'Lembar'], ['Ambon'], ['Kendari'], ['Mamuju', 'Palu'], ['Sorong']], 7: [['Mamuju', 'Palu'], ['Bitung'], ['Oba'], ['Ambon'], ['Celukan Bawang', 'Lembar'], ['Sorong'], ['Kendari']], 8: [['Mamuju'], ['Bitung'], ['Oba'], ['Ambon'], ['Celukan Bawang', 'Lembar'], ['Sorong'], ['Kendari'], ['Palu']], 9: [['Mamuju'], ['Bitung'], ['Oba'], ['Ambon'], ['Lembar'], ['Sorong'], ['Kendari'], ['Palu'], ['Celukan Bawang']]}


def _assert_kmeans_invariance(k, clusters):
    rec = RECORDED_KMEANS_CLUSTERS.get(k)
    if rec is None:
        raise SystemExit(f"[FATAL] no recorded K-Means clusters for k={k}")
    if {frozenset(c) for c in clusters} != {frozenset(c) for c in rec}:
        raise SystemExit(f"[FATAL] K-Means membership for k={k} differs from the recorded "
                         f"production run: {clusters} vs {rec}")
    same_order = [list(c) for c in clusters] == [list(c) for c in rec]
    print(f"  K-Means membership identical to recorded production run "
          f"(list order identical: {same_order})")
    return same_order


def _stage2_manifest(k, clusters, cluster_routes, precomp, sol, same_order, out_dir):
    routes = []
    for c, route in enumerate(cluster_routes):
        legs = ['Biringkassi'] + list(route) + ['Biringkassi']
        routes.append({
            'cluster': c + 1,
            'members': list(clusters[c]),
            'route': legs,
            'route_nm_sea': round(sum(sea_nm(legs[i], legs[i + 1]) for i in range(len(legs) - 1)), 2),
        })
    params = []
    for c in range(len(clusters)):
        for v in range(N_SHIPS):
            pc = precomp[c].get(v)
            if not pc:
                continue
            rec = {'cluster': c + 1, 'vessel': ships[v]['name'], 'feasible': pc.get('feasible')}
            if pc.get('feasible'):
                rec.update({k2: pc[k2] for k2 in ('total_sail_dist', 'RTD_days', 'RTD_periods',
                                                   'arrival_offsets', 'load_days', 'unload_times',
                                                   'sea_days', 'port_days', 'fuel_per_voy',
                                                   'port_charges', 'labuh', 'pandu', 'tunda',
                                                   'tambat', 'tch_annual')})
            else:
                rec['reason'] = pc.get('reason')
            params.append(rec)
    solve = None
    if sol is not None:
        solve = {'status': sol.get('status'), 'objective_incumbent': sol.get('objective'),
                 'objective_bound': sol.get('obj_bound'), 'mip_gap': sol.get('gap'),
                 'gurobi_runtime_sec': sol.get('gurobi_runtime_sec'),
                 'wall_time_sec': sol.get('wall_time_sec'), 'sol_count': sol.get('sol_count'),
                 'costs': sol.get('costs')}
    return {
        'stage': 'Stage 2 (sea-distance version)',
        'k': k,
        'created_at_utc': _dt.datetime.now(_dt.timezone.utc).isoformat(timespec='seconds'),
        'script': os.path.basename(__file__),
        'script_sha256': _sha256_file(os.path.abspath(__file__)),
        'distance_csv': os.path.relpath(SEA_DISTANCE_CSV, _SCRIPT_DIR),
        'distance_csv_sha256': SEA_CSV_SHA256,
        'distance_source_sha256': SEA_META.get('source_sha256'),
        'distance_triangle_exceedances_logged_not_corrected':
            SEA_META.get('triangle_inequality_exceedances_logged_not_corrected'),
        'environment': _environment_record(),
        'kmeans': {'seed': RANDOM_SEED, 'n_init': 20, 'max_iter': 100,
                   'clusters': [list(c) for c in clusters],
                   'coordinates_supplied': HAS_PORT_COORDINATES,
                   'identical_to_recorded_production_run': True,
                   'list_order_identical_to_recorded': same_order},
        'routes': routes,
        'vessel_cluster_parameters': params,
        'constants': {'PERIOD_DAYS': PERIOD_DAYS, 'TOTAL_DAYS': TOTAL_DAYS, 'T_MAX': T_MAX,
                      'COMMISSION_DAYS': COMMISSION_DAYS, 'KURS': KURS, 'HARGA_IFO': HARGA_IFO,
                      'HARGA_MDO': HARGA_MDO, 'CHC_RATE': CHC_RATE,
                      'STOCKOUT_PENALTY': STOCKOUT_PENALTY, 'BUFFER_TIME_PCT': BUFFER_TIME_PCT,
                      'RANDOM_SEED': RANDOM_SEED, 'TIME_LIMIT_SEC': TIME_LIMIT_SEC,
                      'MIPGap': 0.02, 'N_SHIPS': N_SHIPS},
        'solve': solve,
        'output_dir': os.path.relpath(out_dir, _SCRIPT_DIR),
    }


def main():
    # Results are written to OptionC_k{K}_SeaDist_Results/ (never overwritten).
    out_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                           f'OptionC_k{K_CLUSTERS}_SeaDist_Results')
    if os.path.isdir(out_dir) and os.listdir(out_dir):
        raise SystemExit(f"[FATAL] result directory exists and is not empty, refusing to "
                         f"overwrite: {out_dir}")

    print("=" * 70)
    print("MIRP: Stage 1 and Stage 2 (time-charter fleet sizing and inventory routing)")
    print(f"Fleet: {N_SHIPS} existing vessels | 9 destinations | k={K_CLUSTERS} clusters")
    print(f"Time discretization: {PERIOD_DAYS}-day periods (T={T_MAX})")
    print("=" * 70)

    # ── Stage 1a: K-Means ──
    print("\nSTAGE 1a: K-MEANS CLUSTERING")
    print("-" * 40)

    if HAS_PORT_COORDINATES:
        clusters, sse = kmeans_clustering(dest_port_names, K_CLUSTERS)
        same_order = _assert_kmeans_invariance(K_CLUSTERS, clusters)
    else:
        if K_CLUSTERS not in RECORDED_KMEANS_CLUSTERS:
            raise SystemExit(f"[FATAL] no recorded Stage-1 clusters for k={K_CLUSTERS} and "
                             f"no data/port_coordinates.csv to run K-Means")
        clusters = [list(c) for c in RECORDED_KMEANS_CLUSTERS[K_CLUSTERS]]
        same_order = True
        print("  data/port_coordinates.csv not found: using the recorded Stage-1 K-Means "
              "memberships of the reported runs (identical clusters and order).")
    for c, cl_ports in enumerate(clusters):
        D_c = sum(ports[p]['demand'] for p in cl_ports)
        print(f"  Cluster {c+1}: {', '.join(cl_ports)} (demand={D_c:,} ton)")

    # ── Stage 1b: TSP ──
    print("\nSTAGE 1b: TSP ROUTING")
    print("-" * 40)

    cluster_routes = []
    for c, cl_ports in enumerate(clusters):
        route = solve_tsp(cl_ports)
        cluster_routes.append(route)

        dist = 0
        prev_name = 'Biringkassi'
        for p_name in route:
            dist += sea_nm(prev_name, p_name)
            prev_name = p_name
        dist += sea_nm(prev_name, 'Biringkassi')
        route_str = " -> ".join(["Biringkassi"] + route + ["Biringkassi"])
        print(f"  Cluster {c+1}: {route_str} ({dist:.0f} NM)")

    # ── Pre-computation ──
    print("\nPRE-COMPUTATION")
    print("-" * 40)

    precomp = [{} for _ in clusters]
    for c, cl_ports in enumerate(clusters):
        route = cluster_routes[c]
        fc = 0
        for v in range(N_SHIPS):
            pc = precompute_vessel_cluster(ships[v], cl_ports, route)
            precomp[c][v] = pc
            if pc['feasible']:
                fc += 1
                print(f"    {ships[v]['name']:>22} -> Cl{c+1}: "
                      f"RTD={pc['RTD_days']:.1f}d ({pc['RTD_periods']}per), "
                      f"maxFreq={pc['max_freq']}")
        print(f"  Cluster {c+1}: {fc}/{N_SHIPS} feasible vessels")

    # ── Stage 2: Integrated MILP ──
    print("\n" + "=" * 70)
    print("STAGE 2: INTEGRATED MILP")
    print("=" * 70)

    # Build feasible dict
    feasible = {}
    for v in range(N_SHIPS):
        for c in range(len(clusters)):
            pc = precomp[c].get(v)
            if pc and pc['feasible']:
                feasible[(v, c)] = pc

    sol = build_and_solve(clusters, cluster_routes, precomp)

    if sol is None:
        print("\nOPTIMIZATION FAILED")
        os.makedirs(out_dir, exist_ok=True)
        _write_manifest(os.path.join(out_dir, f'run_manifest_k{K_CLUSTERS}.json'),
                        _stage2_manifest(K_CLUSTERS, clusters, cluster_routes, precomp,
                                         None, same_order, out_dir))
        return

    # ── Generate outputs ──
    print("\n" + "=" * 70)
    print("RESULTS")
    print("=" * 70)

    grand = generate_report(sol, clusters, cluster_routes, precomp, feasible, out_dir)

    print(f"\n{'=' * 70}")
    print(f"GRAND TOTAL: Rp {grand:,.0f}")
    print(f"Grand Total: Rp {grand/1e9:,.3f} Billion")
    n_used = len(set(v for v, c in sol['assignments']))
    print(f"Vessels used: {n_used}/{N_SHIPS}")
    print(f"{'=' * 70}")

    _write_manifest(os.path.join(out_dir, f'run_manifest_k{K_CLUSTERS}.json'),
                    _stage2_manifest(K_CLUSTERS, clusters, cluster_routes, precomp,
                                     sol, same_order, out_dir))
    print(f"[Manifest] run manifest written: run_manifest_k{K_CLUSTERS}.json")


if __name__ == '__main__':
    main()
